import io
import openpyxl


def _num(v, default=0):
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", ".").strip())
    except Exception:
        return default


def _int(v, default=0):
    n = _num(v, default)
    try:
        return int(round(n))
    except Exception:
        return default


def _pick_sheet(wb):
    for name in wb.sheetnames:
        if "RE " in name.upper() or name.upper().startswith("RE"):
            ws = wb[name]
            if _has_kecamatan(ws):
                return ws
    for name in wb.sheetnames:
        ws = wb[name]
        if _has_kecamatan(ws):
            return ws
    return None


def _has_kecamatan(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i > 10:
            break
        joined = " ".join(str(c).upper() for c in row if c)
        if "KECAMATAN" in joined and "RUMAH TANGGA" in joined:
            return True
    return False


def parse_re_workbook(content: bytes):
    """Parse a 'RE MURUNG RAYA' formatted Excel. Returns (rows, summary)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = _pick_sheet(wb)
    if ws is None:
        raise ValueError("Sheet rekap kecamatan tidak ditemukan. Pastikan format sesuai template.")

    rows = []
    errors = []
    seen = set()
    no_counter = 0
    for row in ws.iter_rows(values_only=True):
        cells = list(row) + [None] * (24 - len(row)) if len(row) < 24 else list(row)
        idx1, idx2 = cells[1], cells[2]
        # data row: col B integer number and col C non-empty kecamatan name
        if not isinstance(idx1, (int, float)):
            continue
        nama = (str(idx2).strip() if idx2 else "")
        if not nama or nama.upper() in ("JUMLAH", "JUMLAH TOTAL", "KECAMATAN"):
            continue
        no_counter += 1
        rec = {
            "no": no_counter,
            "nama": nama.upper(),
            "jumlah_kelurahan": _int(cells[3]),
            "jumlah_desa": _int(cells[4]),
            "total_rt": _int(cells[5]),
            "keldesa_pln": _int(cells[6]),
            "keldesa_nonpln": _int(cells[7]),
            "keldesa_belum": _int(cells[9]),
            "rt_pln": _int(cells[10]),
            "rt_nonpln": _int(cells[11]),
            "rt_belum": _int(cells[13]),
        }
        rec["jumlah_desa_kel"] = rec["jumlah_kelurahan"] + rec["jumlah_desa"]
        rec["rt_berlistrik"] = rec["rt_pln"] + rec["rt_nonpln"]
        # ratios: prefer file values else compute
        re_pln = _num(cells[14], None)
        re_non = _num(cells[16], None)
        re_tot = _num(cells[18], None)
        rdesa = _num(cells[20], None)
        tot = rec["total_rt"] or 0
        rec["re_pln"] = round(re_pln if re_pln is not None else (rec["rt_pln"] / tot * 100 if tot else 0), 2)
        rec["re_nonpln"] = round(re_non if re_non is not None else (rec["rt_nonpln"] / tot * 100 if tot else 0), 2)
        rec["re_total"] = round(re_tot if re_tot is not None else (rec["rt_berlistrik"] / tot * 100 if tot else 0), 2)
        kd_jumlah = rec["keldesa_pln"] + rec["keldesa_nonpln"] + rec["keldesa_belum"]
        rec["rasio_desa"] = round(rdesa if rdesa is not None else ((rec["keldesa_pln"] + rec["keldesa_nonpln"]) / kd_jumlah * 100 if kd_jumlah else 0), 2)

        # validation
        row_errors = []
        if rec["total_rt"] <= 0:
            row_errors.append("Total rumah tangga kosong / nol")
        if rec["rt_berlistrik"] > rec["total_rt"] and rec["total_rt"] > 0:
            row_errors.append("RT berlistrik melebihi total RT")
        key = rec["nama"]
        if key in seen:
            row_errors.append("Kecamatan duplikat")
        seen.add(key)

        rec["valid"] = len(row_errors) == 0
        rec["errors"] = row_errors
        if row_errors:
            errors.append({"nama": rec["nama"], "errors": row_errors})
        rows.append(rec)

    if not rows:
        raise ValueError("Tidak ada baris data kecamatan yang terbaca dari file.")

    summary = {
        "total_rows": len(rows),
        "valid_rows": sum(1 for r in rows if r["valid"]),
        "error_rows": len(errors),
        "errors": errors,
        "sheet_name": ws.title,
        "totals": {
            "total_rt": sum(r["total_rt"] for r in rows),
            "rt_pln": sum(r["rt_pln"] for r in rows),
            "rt_nonpln": sum(r["rt_nonpln"] for r in rows),
            "rt_belum": sum(r["rt_belum"] for r in rows),
        },
    }
    return rows, summary
